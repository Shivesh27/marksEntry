"""Local API for the Marks Entry Electron application."""
from __future__ import annotations

import argparse
import json
from copy import copy
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

STUDENT_HEADERS = {"name of the student", "student name", "student"}
FATHER_HEADERS = {"father's name", "father name", "father"}
TOTAL_HEADERS = {"total marks", "total"}
PERCENTAGE_HEADERS = {"percentage %", "percentage", "%"}
KNOWN_HEADERS = {"sl.no.", "sl no", "s.no", "s.no.", "admn no.", "admission no", *STUDENT_HEADERS, *FATHER_HEADERS, *TOTAL_HEADERS, *PERCENTAGE_HEADERS}
STATE = {"workbook": None, "path": None, "sheet": None, "layout": None}

def normalise(value): return str(value or "").strip().lower()

def find_layout(sheet):
    for row in range(1, min(sheet.max_row, 50) + 1):
        headers = {normalise(sheet.cell(row, col).value): col for col in range(1, sheet.max_column + 1)}
        student = next((headers[h] for h in STUDENT_HEADERS if h in headers), None)
        father = next((headers[h] for h in FATHER_HEADERS if h in headers), None)
        if student and father:
            return {"header_row": row, "first_data_row": row + 1, "student_col": student, "father_col": father,
                    "subject_cols": [c for c in range(1, sheet.max_column + 1) if sheet.cell(row, c).value not in (None, "") and normalise(sheet.cell(row, c).value) not in KNOWN_HEADERS],
                    "total_col": next((headers[h] for h in TOTAL_HEADERS if h in headers), None),
                    "percentage_col": next((headers[h] for h in PERCENTAGE_HEADERS if h in headers), None)}
    raise ValueError("Could not find the student-name and father-name headings.")

def select_sheet(name):
    sheet = STATE["workbook"][name]
    layout = find_layout(sheet)
    if not layout["subject_cols"]: raise ValueError("No subject columns found on this sheet.")
    STATE.update(sheet=sheet, layout=layout)

def students():
    result, blanks = [], 0
    layout, sheet = STATE["layout"], STATE["sheet"]
    for row in range(layout["first_data_row"], sheet.max_row + 1):
        name = sheet.cell(row, layout["student_col"]).value
        if name not in (None, ""):
            item = {"row": row, "name": str(name), "father": str(sheet.cell(row, layout["father_col"]).value or "")}
            for column in layout["subject_cols"]:
                item[column] = sheet.cell(row, column).value
            result.append(item)
            blanks = 0
        elif result:
            blanks += 1
            if blanks >= 25: break
    return result

def sheet_data():
    layout, sheet = STATE["layout"], STATE["sheet"]
    return {"sheet": sheet.title, "subjects": [{"column": col, "name": str(sheet.cell(layout["header_row"], col).value).strip()} for col in layout["subject_cols"]], "students": students()}

class Handler(BaseHTTPRequestHandler):
    def log_message(self, *_): pass
    def send_json(self, status, data):
        payload = json.dumps(data).encode()
        self.send_response(status); self.send_header("Content-Type", "application/json"); self.send_header("Access-Control-Allow-Origin", "*"); self.send_header("Content-Length", str(len(payload))); self.end_headers(); self.wfile.write(payload)
    def do_OPTIONS(self): self.send_json(204, {})
    def do_GET(self):
        if self.path == "/health": return self.send_json(200, {"ok": True})
        if self.path == "/api/sheets" and STATE["workbook"]: return self.send_json(200, {"sheets": STATE["workbook"].sheetnames, "selected": STATE["sheet"].title})
        if self.path == "/api/data" and STATE["sheet"]: return self.send_json(200, sheet_data())
        self.send_json(404, {"error": "Not found"})
    def do_POST(self):
        try:
            length = int(self.headers.get("Content-Length", 0)); body = json.loads(self.rfile.read(length) or b"{}")
            if self.path == "/api/open":
                STATE["workbook"] = load_workbook(body["path"]); STATE["path"] = body["path"]; select_sheet(STATE["workbook"].sheetnames[0])
                return self.send_json(200, {"sheets": STATE["workbook"].sheetnames, "selected": STATE["sheet"].title})
            if self.path == "/api/select-sheet": select_sheet(body["sheet"]); return self.send_json(200, sheet_data())
            if self.path == "/api/marks":
                layout, sheet, row = STATE["layout"], STATE["sheet"], int(body["row"])
                marks = [float(value) for value in body["marks"]]
                if len(marks) != len(layout["subject_cols"]) or any(mark < 0 or mark > 100 for mark in marks): raise ValueError("Each subject mark must be between 0 and 100.")
                for column, mark in zip(layout["subject_cols"], marks): sheet.cell(row, column).value = int(mark) if mark.is_integer() else mark
                total = sum(marks)
                if layout["total_col"]: sheet.cell(row, layout["total_col"]).value = total
                if layout["percentage_col"]: sheet.cell(row, layout["percentage_col"]).value = round(total / len(marks), 2)
                return self.send_json(200, {"total": total, "percentage": round(total / len(marks), 2)})
            if self.path == "/api/save": STATE["workbook"].save(STATE["path"]); return self.send_json(200, {"saved": True})
            if self.path == "/api/source-sheets":
                source = load_workbook(body["path"], read_only=True, data_only=True)
                return self.send_json(200, {"sheets": source.sheetnames})
            if self.path == "/api/import": return self.import_students(body)
            self.send_json(404, {"error": "Not found"})
        except Exception as exc: self.send_json(400, {"error": str(exc)})
    def import_students(self, body):
        destination = body["newSheet"].strip()
        if not destination or destination in STATE["workbook"].sheetnames: raise ValueError("Choose a unique new sheet name.")
        source = load_workbook(body["sourcePath"], read_only=True, data_only=True)[body["sourceSheet"]]
        name_col, father_col, start = column_index_from_string(body["nameColumn"]), column_index_from_string(body["fatherColumn"]), int(body["firstRow"])
        imported, blanks = [], 0
        for row in range(start, source.max_row + 1):
            name = source.cell(row, name_col).value
            if name not in (None, ""): imported.append((str(name).strip(), str(source.cell(row, father_col).value or "").strip())); blanks = 0
            elif imported:
                blanks += 1
                if blanks >= 25: break
        if not imported: raise ValueError("No students found for those import settings.")
        new = STATE["workbook"].copy_worksheet(STATE["sheet"]); new.title = destination; layout = find_layout(new); style_row = layout["first_data_row"]
        for row in range(style_row, max(new.max_row, style_row + len(imported)) + 1):
            for col in range(1, new.max_column + 1): new.cell(row, col).value = None
        for number, (name, father) in enumerate(imported, 1):
            row = style_row + number - 1
            if row != style_row:
                for col in range(1, new.max_column + 1): new.cell(row, col)._style = copy(new.cell(style_row, col)._style)
            new.cell(row, 1).value, new.cell(row, layout["student_col"]).value, new.cell(row, layout["father_col"]).value = number, name, father
        select_sheet(destination); STATE["workbook"].save(STATE["path"])
        self.send_json(200, {"count": len(imported), "data": sheet_data()})

if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--port", type=int, default=8765); args = parser.parse_args()
    ThreadingHTTPServer(("127.0.0.1", args.port), Handler).serve_forever()
